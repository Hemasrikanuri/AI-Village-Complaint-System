import io
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.database import get_db
from app.models import Complaint, ComplaintStatus, PriorityLevel, UserRole, User
from app.auth import require_roles

router = APIRouter(prefix="/reports", tags=["Reports"], dependencies=[Depends(require_roles([UserRole.ADMIN]))])

@router.get("/excel")
def export_complaints_excel(
    status_filter: Optional[ComplaintStatus] = Query(None),
    priority_filter: Optional[PriorityLevel] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(Complaint)
    if status_filter:
        query = query.filter(Complaint.status == status_filter)
    if priority_filter:
        query = query.filter(Complaint.priority == priority_filter)
        
    complaints = query.order_by(Complaint.created_at.desc()).all()

    # Create OpenPyXL Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GramSetu Complaints"

    # Header title
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"GramSetu - Village Complaint Report ({datetime.utcnow().strftime('%Y-%m-%d %H:%M')})"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    headers = [
        "Ref ID", "Title", "Citizen", "Village", "Category", "Department",
        "Priority", "Status", "Assigned Officer", "AI Confidence", "Date"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[2].height = 25

    # Rows
    for row_idx, c in enumerate(complaints, start=3):
        row_data = [
            c.reference_id,
            c.title,
            c.citizen.name if c.citizen else "N/A",
            c.village.name if c.village else "N/A",
            c.category.name if c.category else "N/A",
            c.category.department.name if c.category and c.category.department else "N/A",
            c.priority.value,
            c.status.value,
            c.assigned_officer.name if c.assigned_officer else "Unassigned",
            f"{int(c.ai_confidence * 100)}%",
            c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else ""
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = center_align if col_idx in [1, 7, 8, 10, 11] else left_align
            cell.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"GramSetu_Complaints_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/pdf")
def export_complaints_pdf(
    status_filter: Optional[ComplaintStatus] = Query(None),
    priority_filter: Optional[PriorityLevel] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(Complaint)
    if status_filter:
        query = query.filter(Complaint.status == status_filter)
    if priority_filter:
        query = query.filter(Complaint.priority == priority_filter)
        
    complaints = query.order_by(Complaint.created_at.desc()).limit(100).all()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2563EB'),
        spaceAfter=8
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#4B5563'),
        spaceAfter=15
    )
    table_text = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontSize=8,
        leading=10
    )
    header_text = ParagraphStyle(
        'HeaderText',
        parent=styles['Normal'],
        fontSize=9,
        leading=11,
        textColor=colors.white,
        fontName='Helvetica-Bold'
    )

    elements = []
    elements.append(Paragraph("GramSetu - Panchayat Complaint Summary", title_style))
    elements.append(Paragraph(f"Generated on {datetime.utcnow().strftime('%B %d, %Y at %H:%M UTC')} | Total Records: {len(complaints)}", subtitle_style))
    elements.append(Spacer(1, 10))

    # Table headers
    data = [[
        Paragraph("Ref ID", header_text),
        Paragraph("Title", header_text),
        Paragraph("Village", header_text),
        Paragraph("Category", header_text),
        Paragraph("Priority", header_text),
        Paragraph("Status", header_text),
        Paragraph("Assigned Officer", header_text)
    ]]

    for c in complaints:
        data.append([
            Paragraph(c.reference_id, table_text),
            Paragraph(c.title[:30] + ('...' if len(c.title) > 30 else ''), table_text),
            Paragraph(c.village.name if c.village else "-", table_text),
            Paragraph(c.category.name if c.category else "-", table_text),
            Paragraph(c.priority.value, table_text),
            Paragraph(c.status.value, table_text),
            Paragraph(c.assigned_officer.name if c.assigned_officer else "Unassigned", table_text)
        ])

    table = Table(data, colWidths=[65, 120, 70, 80, 55, 65, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
    ]))

    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    filename = f"GramSetu_Summary_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
